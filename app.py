from __future__ import annotations

import io
import json
import textwrap
from collections import Counter
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook, load_workbook

EXCEL_PATH = Path(__file__).resolve().parent / "BUGS Reportados SIR 2024.xlsx"
HOST = "127.0.0.1"
PORT = 8000

CANDIDATE_FILTERS = [
    "anio",
  "sprint Reportes",
  "Sprint levantado",
  "sprint",
  "Sprint Solución",
    "Sprint despliegue",
    "PAIS",
    "País",
    "Status",
    "Tipo",
    "Responsable de Reporte",
    "prioridad DB",
]

CLOSED_STATUS_VALUES = {
    "done",
    "solventado",
    "cerrado",
    "cerrada",
    "completado",
    "finalizado",
    "closed",
    "resuelto",
}

WORKBOOK_DATA: dict[str, dict[str, object]] = {}
DEFAULT_SHEET = ""
LAST_REFRESH = ""

HTML_PAGE = """<!doctype html>
<html lang="es">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Tablero Gerencial de Defectos</title>
  <style>
    :root {
      --kfc-red: #e4002b;
      --kfc-black: #131313;
      --kfc-cream: #f6f2ea;
      --kfc-white: #ffffff;
      --line: #e4ddd3;
      --text: #1a1a1a;
      --muted: #666056;
      --ok: #2e7d32;
      --warning: #b16a00;
      --shadow: 0 16px 42px rgba(0, 0, 0, 0.12);
      --radius: 18px;
    }

    * { box-sizing: border-box; }

    html, body {
      overflow-x: hidden;
    }

    body {
      margin: 0;
      font-family: "Segoe UI", Tahoma, Geneva, Verdana, sans-serif;
      color: var(--text);
      background:
        radial-gradient(circle at 2% 0%, rgba(228, 0, 43, 0.15), transparent 22%),
        radial-gradient(circle at 98% 0%, rgba(19, 19, 19, 0.1), transparent 26%),
        var(--kfc-cream);
      min-height: 100vh;
    }

    .layout {
      max-width: 1700px;
      margin: 0 auto;
      display: grid;
      grid-template-columns: 290px 1fr;
      gap: 18px;
      padding: 18px;
    }

    .layout,
    .main,
    .panel {
      min-width: 0;
    }

    .panel {
      background: var(--kfc-white);
      border: 1px solid var(--line);
      border-radius: var(--radius);
      box-shadow: var(--shadow);
    }

    .options {
      position: sticky;
      top: 18px;
      align-self: start;
      padding: 16px;
      background:
        linear-gradient(180deg, #ffffff 0%, #fbf8f2 100%);
      border-top: 6px solid var(--kfc-red);
    }

    .options h2 {
      margin: 0;
      font-size: 20px;
      letter-spacing: -0.02em;
    }

    .options p {
      margin: 6px 0 14px;
      color: var(--muted);
      line-height: 1.4;
      font-size: 13px;
    }

    .option-stack {
      display: grid;
      gap: 10px;
    }

    .option-btn {
      border: 1px solid var(--kfc-red);
      border-radius: 12px;
      background: var(--kfc-red);
      color: #fff;
      font-weight: 700;
      padding: 10px 12px;
      cursor: pointer;
      transition: transform 120ms ease, filter 120ms ease;
      text-align: left;
    }

    .option-btn:hover {
      transform: translateY(-1px);
      filter: brightness(0.95);
    }

    .option-btn.secondary {
      background: #fff;
      color: var(--kfc-black);
      border-color: #cec6bb;
    }

    .option-btn.secondary:hover {
      filter: none;
      border-color: var(--kfc-black);
    }

    .status-box {
      margin-top: 12px;
      padding: 10px;
      border-radius: 10px;
      border: 1px dashed #d0c6b9;
      background: #faf7f2;
      color: var(--muted);
      font-size: 12px;
      line-height: 1.45;
    }

    .main {
      display: grid;
      gap: 16px;
    }

    .hero {
      padding: 22px;
      background:
        linear-gradient(130deg, #1a1a1a 0%, #212121 45%, #2a2a2a 100%);
      color: #fff;
      border: 1px solid rgba(255, 255, 255, 0.08);
      border-radius: var(--radius);
      box-shadow: var(--shadow);
      position: relative;
      overflow: hidden;
    }

    .hero::after {
      content: "";
      position: absolute;
      right: -60px;
      top: -50px;
      width: 220px;
      height: 220px;
      border-radius: 50%;
      background: radial-gradient(circle, rgba(228, 0, 43, 0.35), transparent 68%);
    }

    .eyebrow {
      color: #ffc6d2;
      text-transform: uppercase;
      letter-spacing: 0.16em;
      font-size: 11px;
      font-weight: 700;
    }

    h1 {
      margin: 8px 0 8px;
      font-size: clamp(30px, 4vw, 44px);
      line-height: 1.02;
      letter-spacing: -0.03em;
    }

    .subtitle {
      margin: 0;
      max-width: 1000px;
      color: #ebe7de;
      line-height: 1.55;
      font-size: 15px;
    }

    .pill-row {
      display: flex;
      gap: 8px;
      flex-wrap: wrap;
      margin-top: 14px;
    }

    .pill {
      display: inline-flex;
      align-items: center;
      gap: 6px;
      background: rgba(255, 255, 255, 0.08);
      border: 1px solid rgba(255, 255, 255, 0.18);
      border-radius: 999px;
      padding: 7px 10px;
      font-size: 12px;
      color: #f8f4ef;
    }

    .loading {
      margin-top: 10px;
      color: #efe7da;
      font-size: 13px;
      display: inline-flex;
      align-items: center;
      gap: 8px;
    }

    .spinner {
      width: 14px;
      height: 14px;
      border-radius: 50%;
      border: 2px solid rgba(255, 255, 255, 0.35);
      border-top-color: #fff;
      animation: spin 0.8s linear infinite;
    }

    @keyframes spin { to { transform: rotate(360deg); } }

    .kpis {
      display: grid;
      gap: 12px;
      grid-template-columns: repeat(auto-fit, minmax(210px, 1fr));
    }

    .kpi {
      padding: 14px;
      border-left: 5px solid var(--kfc-red);
    }

    .kpi .label {
      color: var(--muted);
      text-transform: uppercase;
      letter-spacing: 0.1em;
      font-size: 11px;
      margin-bottom: 8px;
      font-weight: 700;
    }

    .kpi .value {
      font-size: 28px;
      font-weight: 800;
      letter-spacing: -0.03em;
      color: var(--kfc-black);
    }

    .kpi .hint {
      margin-top: 8px;
      color: var(--muted);
      font-size: 12px;
    }

    .section {
      padding: 16px;
    }

    .section-head {
      margin-bottom: 12px;
    }

    .section-title {
      margin: 0;
      font-size: 19px;
    }

    .section-subtitle {
      margin: 4px 0 0;
      color: var(--muted);
      font-size: 13px;
    }

    .filters {
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
      gap: 10px;
      align-items: end;
    }

    .field {
      display: flex;
      flex-direction: column;
      gap: 6px;
    }

    .field label {
      font-size: 11px;
      text-transform: uppercase;
      color: var(--muted);
      letter-spacing: 0.09em;
      font-weight: 700;
    }

    .field input, .field select {
      width: 100%;
      border: 1px solid #d4ccbe;
      border-radius: 11px;
      padding: 10px;
      background: #fff;
      color: var(--text);
      outline: none;
    }

    .field select[multiple] {
      min-height: 130px;
      padding: 6px;
    }

    .combo {
      position: relative;
    }

    .combo-toggle {
      width: 100%;
      border: 1px solid #d4ccbe;
      border-radius: 11px;
      padding: 10px;
      background: #fff;
      color: var(--text);
      cursor: pointer;
      text-align: left;
      display: flex;
      justify-content: space-between;
      align-items: center;
      gap: 8px;
    }

    .combo-toggle:hover {
      border-color: var(--kfc-red);
    }

    .combo-toggle .chev {
      color: var(--kfc-red);
      font-size: 11px;
      flex: 0 0 auto;
    }

    .combo-panel {
      display: none;
      position: absolute;
      z-index: 20;
      top: calc(100% + 4px);
      left: 0;
      right: 0;
      max-height: 240px;
      overflow-y: auto;
      background: #fff;
      border: 1px solid #d4ccbe;
      border-radius: 11px;
      box-shadow: 0 14px 32px rgba(0, 0, 0, 0.16);
      padding: 6px;
    }

    .combo.open .combo-panel {
      display: block;
    }

    .combo-option {
      display: flex;
      align-items: center;
      gap: 8px;
      padding: 7px 8px;
      border-radius: 8px;
      cursor: pointer;
      font-size: 13px;
    }

    .combo-option:hover {
      background: #fbf1f3;
    }

    .combo-option input {
      width: auto;
      accent-color: var(--kfc-red);
    }

    .footer-actions {
      display: flex;
      gap: 10px;
      flex-wrap: wrap;
      align-items: center;
      justify-content: flex-end;
    }

    .footer-actions .option-btn {
      min-width: 180px;
      text-align: center;
    }

    .field-actions {
      display: flex;
      flex-direction: column;
      justify-content: flex-end;
      gap: 6px;
    }

    .field-actions label {
      min-height: 14px;
    }

    .field-actions .option-btn {
      width: 100%;
      text-align: center;
      min-height: 40px;
      font-size: 13px;
      padding: 8px 12px;
    }

    .field-actions .option-btn.secondary {
      border-color: var(--kfc-red);
      color: var(--kfc-red);
      background: #fff5f7;
      font-weight: 700;
      letter-spacing: 0.02em;
    }

    .field-actions .option-btn.secondary:hover {
      background: #ffc6d2;
      color: var(--kfc-black);
    }

    .field .hint {
      color: var(--muted);
      font-size: 11px;
      line-height: 1.3;
    }

    .field input:focus, .field select:focus {
      border-color: var(--kfc-red);
      box-shadow: 0 0 0 2px rgba(228, 0, 43, 0.1);
    }

    .charts {
      display: grid;
      gap: 12px;
      grid-template-columns: repeat(2, minmax(0, 1fr));
    }

    .bar-list {
      display: grid;
      gap: 10px;
      margin-top: 10px;
    }

    .bar-item { display: grid; gap: 6px; }

    .bar-top {
      display: flex;
      justify-content: space-between;
      font-size: 13px;
      color: var(--muted);
    }

    .bar-track {
      height: 10px;
      border-radius: 999px;
      background: #f1ece2;
      position: relative;
      overflow: hidden;
    }

    .bar-fill {
      position: absolute;
      inset: 0 auto 0 0;
      border-radius: inherit;
      background: linear-gradient(90deg, var(--kfc-red), #ff7a91);
    }

    .table-wrap {
      border: 1px solid var(--line);
      border-radius: 14px;
      overflow: auto;
      max-height: none;
      background: #fff;
    }

    table {
      width: 100%;
      border-collapse: collapse;
      table-layout: fixed;
      font-size: 13px;
    }

    thead th {
      position: sticky;
      top: 0;
      background: #1f1f1f;
      color: #fff;
      text-align: left;
      padding: 11px 10px;
      white-space: normal;
      word-break: break-word;
    }

    tbody td {
      text-align: center;
      padding: 12px;
      border-bottom: 1px solid #ede6da;
      vertical-align: middle;
      white-space: normal;
      word-break: break-word;
      overflow-wrap: anywhere;
      line-height: 1.5;
    }

    tbody tr:nth-child(even) td {
      background: #faf7f1;
    }

    .empty {
      text-align: center;
      color: var(--muted);
      padding: 16px;
    }

    .modal-overlay {
      position: fixed;
      inset: 0;
      z-index: 50;
      display: none;
      align-items: center;
      justify-content: center;
      padding: 16px;
      background: rgba(0, 0, 0, 0.52);
      backdrop-filter: blur(5px);
    }

    .modal-overlay.open {
      display: flex;
    }

    .modal-dialog {
      width: min(1300px, 100%);
      max-height: 92vh;
      overflow: auto;
      background: #fff;
      border: 1px solid #d8cec0;
      border-radius: 18px;
      padding: 14px;
      box-shadow: 0 26px 90px rgba(0, 0, 0, 0.28);
    }

    .modal-head {
      display: flex;
      justify-content: space-between;
      align-items: flex-start;
      gap: 12px;
      margin-bottom: 10px;
    }

    .modal-title {
      margin: 0;
      font-size: 23px;
      color: var(--kfc-black);
    }

    .modal-subtitle {
      margin: 4px 0 0;
      color: var(--muted);
      font-size: 13px;
    }

    .close-btn {
      border: 1px solid #c7bcae;
      background: #fff;
      color: #111;
      width: 38px;
      height: 38px;
      border-radius: 999px;
      cursor: pointer;
      font-size: 18px;
    }

    .pie-grid {
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      gap: 12px;
    }

    .pie-card {
      border: 1px solid #e3dbce;
      border-radius: 14px;
      padding: 12px;
      background: #fff;
      display: flex;
      flex-direction: column;
      gap: 10px;
      min-height: 330px;
    }

    .pie-card h3 {
      margin: 0;
      font-size: 16px;
    }

    .pie-card p {
      margin: 4px 0 0;
      color: var(--muted);
      font-size: 12px;
      line-height: 1.45;
    }

    .pie-figure {
      width: min(220px, 100%);
      margin: 0 auto;
      aspect-ratio: 1 / 1;
      border-radius: 50%;
      position: relative;
      background: conic-gradient(#e4002b 0deg 360deg);
    }

    .pie-figure::after {
      content: "";
      position: absolute;
      inset: 24%;
      border-radius: 50%;
      background: #fff;
      border: 1px solid #e7ddcf;
    }

    .pie-center {
      position: absolute;
      inset: 24%;
      z-index: 1;
      display: flex;
      flex-direction: column;
      align-items: center;
      justify-content: center;
      text-align: center;
      pointer-events: none;
    }

    .pie-center strong {
      font-size: 25px;
      color: var(--kfc-black);
    }

    .pie-center span {
      color: var(--muted);
      font-size: 11px;
      max-width: 130px;
      line-height: 1.3;
    }

    .pie-legend {
      display: grid;
      gap: 8px;
      margin-top: auto;
    }

    .legend-item {
      display: flex;
      gap: 8px;
      align-items: flex-start;
      font-size: 13px;
    }

    .legend-swatch {
      width: 11px;
      height: 11px;
      border-radius: 50%;
      margin-top: 4px;
      flex: 0 0 auto;
    }

    .legend-meta {
      color: var(--muted);
      margin-left: auto;
      font-size: 11px;
      white-space: nowrap;
    }

    .modal-actions {
      display: flex;
      gap: 8px;
      align-items: center;
      flex: 0 0 auto;
    }

    .modal-actions .option-btn {
      padding: 8px 14px;
    }


    .page-view {
      display: block;
    }

    .page-view.hidden-section {
      display: none;
    }

    .hidden-section {
      display: none;
    }

    .submenu {
      display: flex;
      flex-direction: column;
      gap: 4px;
      padding: 4px 0 4px 8px;
      border-left: 2px solid var(--kfc-red);
      margin-top: 4px;
    }

    .submenu.hidden {
      display: none;
    }

    .submenu-item {
      font-size: 12px !important;
      padding: 8px !important;
      min-height: 0 !important;
    }

    .section-toggle {
      scroll-behavior: smooth;
    }

    @media print {
      body * { visibility: hidden; }
      #insightModal, #insightModal * { visibility: visible; }
      #pasteleProitySection, #pasteleProitySection * { visibility: visible; }
      #pasteleProitySection {
        position: absolute;
        inset: 0;
        display: block;
        background: #fff;
        padding: 20px;
        visibility: visible;
      }
      #insightModal.open {
        position: absolute;
        inset: 0;
        display: block;
        background: #fff;
        padding: 0;
      }
      .modal-dialog {
        width: 100%;
        max-height: none;
        overflow: visible;
        border: none;
        box-shadow: none;
        border-radius: 0;
      }
      .section-head { display: block; }
      .modal-actions { display: none !important; }
      .pie-grid { page-break-inside: avoid; }
      .pie-card { break-inside: avoid; page-break-inside: avoid; }
      .pie-figure, .legend-swatch {
        -webkit-print-color-adjust: exact;
        print-color-adjust: exact;
      }
    }

    @media (max-width: 1300px) {
      .layout {
        grid-template-columns: 1fr;
      }

      .options {
        position: static;
      }
    }

    @media (max-width: 1080px) {
      .filters,
      .charts,
      .pie-grid {
        grid-template-columns: 1fr;
      }

      .table-wrap {
        max-height: none;
      }
    }

    .pie-expand-btn {
      position: absolute;
      top: 12px;
      right: 12px;
      width: 32px;
      height: 32px;
      border: 1px solid #d4ccbe;
      border-radius: 8px;
      background: #fff;
      color: var(--kfc-red);
      font-size: 16px;
      cursor: pointer;
      display: flex;
      align-items: center;
      justify-content: center;
      transition: all 200ms ease;
      z-index: 10;
    }

    .pie-expand-btn:hover {
      background: #fff5f7;
      border-color: var(--kfc-red);
      transform: scale(1.1);
    }

    .pie-card {
      position: relative;
    }

    .pie-modal-overlay {
      position: fixed;
      inset: 0;
      z-index: 60;
      display: none;
      align-items: center;
      justify-content: center;
      padding: 20px;
      background: rgba(0, 0, 0, 0.55);
      backdrop-filter: blur(6px);
    }

    .pie-modal-overlay.open {
      display: flex;
    }

    .pie-modal-content {
      width: min(800px, 100%);
      max-height: 85vh;
      overflow: auto;
      background: #fff;
      border: 1px solid #d8cec0;
      border-radius: 18px;
      padding: 24px;
      box-shadow: 0 26px 90px rgba(0, 0, 0, 0.3);
      position: relative;
    }

    .pie-modal-close {
      position: absolute;
      top: 16px;
      right: 16px;
      width: 40px;
      height: 40px;
      border: 1px solid #c7bcae;
      background: #fff;
      color: #111;
      border-radius: 50%;
      cursor: pointer;
      font-size: 22px;
      display: flex;
      align-items: center;
      justify-content: center;
      transition: all 200ms ease;
    }

    .pie-modal-close:hover {
      background: #f5f5f5;
      border-color: #111;
    }

    .pie-modal-figure {
      width: 100%;
      max-width: 500px;
      aspect-ratio: 1 / 1;
      margin: 0 auto;
      border-radius: 50%;
      position: relative;
    }

    .pie-modal-title {
      margin: 0 0 8px;
      font-size: 22px;
      color: var(--kfc-black);
    }

    .pie-modal-subtitle {
      margin: 0 0 20px;
      color: var(--muted);
      font-size: 13px;
    }

    .pie-modal-legend {
      display: grid;
      gap: 10px;
      margin-top: 20px;
    }
  </style>
</head>
<body>
  <div class="layout">
    <aside class="panel options">
      <h2>Opciones</h2>
      <p>Módulo lateral para acciones gerenciales rápidas.</p>
      <div class="option-stack">
        <button class="option-btn" id="navTableroGeneral" type="button">Tablero General</button>

        <div style="margin-top: 8px;">
          <button class="option-btn secondary" id="navPastelesToggle" type="button" style="width:100%;text-align:left;display:flex;justify-content:space-between;align-items:center;">
            <span>Pasteles</span>
            <span id="pastelesChev" style="font-size:10px;">▶</span>
          </button>
          <div id="pastelesSubmenu" class="submenu hidden">
            <button class="option-btn secondary submenu-item" id="navPastelesSprint" type="button">Pasteles Sprint</button>
            <button class="option-btn secondary submenu-item" id="navPastelesPrority" type="button">Pasteles Prioridad DB</button>
          </div>
        </div>

        <button class="option-btn" id="refreshExcelBtn" type="button">Actualizar desde Excel</button>
        <button class="option-btn secondary" id="exportExcelBtn" type="button">Exportar Excel</button>
        <button class="option-btn secondary" id="exportPdfBtn" type="button">Exportar PDF</button>
      </div>
      <div class="status-box" id="optionStatus">Sincronizado con el archivo actual.</div>
      <div style="margin-top: 16px; padding-top: 12px; border-top: 1px solid #ddd; text-align: center; font-size: 11px; color: #999;">
        <strong>Tablero Gerencial v1.0.0</strong><br/>
        <span>© 2026 - Todos los derechos reservados</span>
      </div>
    </aside>

    <main class="main">
      <div id="pageTableroGeneral" class="page-view">
        <section class="hero">
          <div class="eyebrow">Tablero gerencial</div>
          <h1>Defectos y operación por hoja</h1>
          <p class="subtitle">
            Diseño con paleta KFC. La hoja del Excel es el filtro principal y todos los indicadores,
            pasteles y exportaciones se recalculan con la vista filtrada actual.
          </p>
          <div style="margin-top:10px;">
            <button class="option-btn" id="refreshExcelBtnMain" type="button">Actualizar Excel (principal)</button>
          </div>
          <div class="pill-row" id="summaryPills"></div>
          <div class="loading" id="loadingState"><span class="spinner"></span> Cargando datos</div>
        </section>

        <section class="kpis" id="kpiGrid"></section>

        <section class="panel section">
          <div class="section-head">
            <h2 class="section-title">Filtros</h2>
            <p class="section-subtitle">La hoja es el filtro más importante. Los demás filtros se adaptan automáticamente.</p>
          </div>
          <div class="filters" id="filters"></div>
        </section>

        <section class="charts">
          <article class="panel section">
            <div class="section-head">
              <h2 class="section-title">Distribución por estado</h2>
              <p class="section-subtitle">Lectura rápida del estado de defectos.</p>
            </div>
            <div id="statusChart" class="bar-list"></div>
          </article>

          <article class="panel section">
            <div class="section-head">
              <h2 class="section-title">Distribución por país</h2>
              <p class="section-subtitle">Países más afectados en la hoja activa.</p>
            </div>
            <div id="countryChart" class="bar-list"></div>
          </article>
        </section>

        <section class="panel section">
          <div class="section-head">
            <h2 class="section-title">Detalle de registros</h2>
            <p class="section-subtitle" id="tableSubtitle">Tabla dinámica filtrada desde el Excel.</p>
          </div>
          <div class="table-wrap" id="tableWrap"></div>
        </section>
      </div>

      <section class="panel section hidden-section" id="pasteleSprintSection">
        <div class="section-head">
          <h2 class="section-title">Pasteles de Sprint</h2>
          <p class="section-subtitle">Distribuciones del conjunto visible con filtros activos.</p>
        </div>
        <div class="pie-grid" id="pasteleSprintGrid">
          <section class="pie-card">
            <button class="pie-expand-btn" data-expand-pie="pieSprint1" title="Expandir gráfico">⛶</button>
            <div><h3>Distribución por estado</h3><p>Participación por estado / status en los registros visibles.</p></div>
            <div class="pie-figure" id="pieSprint1"><div class="pie-center"><strong id="pieSprint1Value">0%</strong><span id="pieSprint1Label">Sin datos</span></div></div>
            <div class="pie-legend" id="pieSprint1Legend"></div>
          </section>

          <section class="pie-card">
            <button class="pie-expand-btn" data-expand-pie="pieSprint2" title="Expandir gráfico">⛶</button>
            <div><h3>Porcentaje de defecto por país</h3><p>Participación por país en los registros visibles.</p></div>
            <div class="pie-figure" id="pieSprint2"><div class="pie-center"><strong id="pieSprint2Value">0%</strong><span id="pieSprint2Label">Sin datos</span></div></div>
            <div class="pie-legend" id="pieSprint2Legend"></div>
          </section>

          <section class="pie-card">
            <button class="pie-expand-btn" data-expand-pie="pieSprint3" title="Expandir gráfico">⛶</button>
            <div><h3>Despliegue por sprint</h3><p>Tomado del campo Sprint despliegue.</p></div>
            <div class="pie-figure" id="pieSprint3"><div class="pie-center"><strong id="pieSprint3Value">0%</strong><span id="pieSprint3Label">Sin datos</span></div></div>
            <div class="pie-legend" id="pieSprint3Legend"></div>
          </section>

          <section class="pie-card">
            <button class="pie-expand-btn" data-expand-pie="pieSprint4" title="Expandir gráfico">⛶</button>
            <div><h3>Prioridad DB</h3><p>Distribución por prioridad DB de la hoja actual.</p></div>
            <div class="pie-figure" id="pieSprint4"><div class="pie-center"><strong id="pieSprint4Value">0%</strong><span id="pieSprint4Label">Sin datos</span></div></div>
            <div class="pie-legend" id="pieSprint4Legend"></div>
          </section>

          <section class="pie-card">
            <button class="pie-expand-btn" data-expand-pie="pieSprint5" title="Expandir gráfico">⛶</button>
            <div><h3>Responsable del reporte</h3><p>Participación por responsable de reporte.</p></div>
            <div class="pie-figure" id="pieSprint5"><div class="pie-center"><strong id="pieSprint5Value">0%</strong><span id="pieSprint5Label">Sin datos</span></div></div>
            <div class="pie-legend" id="pieSprint5Legend"></div>
          </section>
        </div>
      </section>

      <section class="panel section hidden-section" id="pasteleProitySection">
        <div class="section-head">
          <h2 class="section-title">Análisis: Prioridad DB por Sprint</h2>
          <p class="section-subtitle" id="prioBySprintHint">Marca uno o más sprints en el filtro para comparar su prioridad DB.</p>
          <div style="margin-top: 12px; display: flex; gap: 8px;">
            <button class="option-btn" id="savePrioDbPdfBtn" type="button" style="font-size: 12px; padding: 8px 12px;">Descargar PDF</button>
            <button class="option-btn secondary" id="savePrioDbExcelBtn" type="button" style="font-size: 12px; padding: 8px 12px;">Descargar Excel</button>
          </div>
        </div>
        <div class="pie-grid" id="prioBySprintGrid"></div>
      </section>

    </main>
  </div>

  <div class="pie-modal-overlay" id="pieExpandModal" aria-hidden="true">
    <div class="pie-modal-content" role="dialog" aria-modal="true" aria-labelledby="pieExpandTitle">
      <button class="pie-modal-close" id="closePieExpandBtn" aria-label="Cerrar gráfico expandido">×</button>
      <h2 class="pie-modal-title" id="pieExpandTitle">Gráfico expandido</h2>
      <p class="pie-modal-subtitle" id="pieExpandSubtitle">Vista detallada del gráfico</p>
      <div class="pie-modal-figure" id="pieExpandFigure"></div>
      <div class="pie-modal-legend" id="pieExpandLegend"></div>
    </div>
  </div>

  <div class="modal-overlay" id="insightModal" aria-hidden="true">
    <div class="modal-dialog" role="dialog" aria-modal="true" aria-labelledby="insightTitle">
      <div class="modal-head">
        <div>
          <h2 class="modal-title" id="insightTitle">Pasteles gerenciales</h2>
          <p class="modal-subtitle">Distribuciones del conjunto visible con filtros activos.</p>
          <p class="modal-subtitle" id="insightSprintLabel"><strong>Sprint:</strong> Todos</p>
        </div>
        <div class="modal-actions">
          <button class="option-btn" id="savePiesPdfBtn" type="button">Guardar PDF</button>
          <button class="close-btn" id="closeInsightsBtn" type="button" aria-label="Cerrar">×</button>
        </div>
      </div>

      <div class="pie-grid">
        <section class="pie-card">
          <button class="pie-expand-btn" data-expand-pie="pieStatus" title="Expandir gráfico">⛶</button>
          <div>
            <h3>Distribución por estado</h3>
            <p>Participación por estado / status en los registros visibles.</p>
          </div>
          <div class="pie-figure" id="pieStatus"><div class="pie-center"><strong id="pieStatusValue">0%</strong><span id="pieStatusLabel">Sin datos</span></div></div>
          <div class="pie-legend" id="pieStatusLegend"></div>
        </section>

        <section class="pie-card">
          <button class="pie-expand-btn" data-expand-pie="pieCountry" title="Expandir gráfico">⛶</button>
          <div>
            <h3>Porcentaje de defecto por país</h3>
            <p>Participación por país en los registros visibles.</p>
          </div>
          <div class="pie-figure" id="pieCountry"><div class="pie-center"><strong id="pieCountryValue">0%</strong><span id="pieCountryLabel">Sin datos</span></div></div>
          <div class="pie-legend" id="pieCountryLegend"></div>
        </section>

        <section class="pie-card">
          <button class="pie-expand-btn" data-expand-pie="pieDeployment" title="Expandir gráfico">⛶</button>
          <div>
            <h3>Despliegue por sprint</h3>
            <p>Tomado del campo Sprint despliegue.</p>
          </div>
          <div class="pie-figure" id="pieDeployment"><div class="pie-center"><strong id="pieDeploymentValue">0%</strong><span id="pieDeploymentLabel">Sin datos</span></div></div>
          <div class="pie-legend" id="pieDeploymentLegend"></div>
        </section>

        <section class="pie-card">
          <button class="pie-expand-btn" data-expand-pie="piePriorityDb" title="Expandir gráfico">⛶</button>
          <div>
            <h3>Prioridad DB</h3>
            <p>Distribución por prioridad DB de la hoja actual.</p>
          </div>
          <div class="pie-figure" id="piePriorityDb"><div class="pie-center"><strong id="piePriorityDbValue">0%</strong><span id="piePriorityDbLabel">Sin datos</span></div></div>
          <div class="pie-legend" id="piePriorityDbLegend"></div>
        </section>

        <section class="pie-card">
          <button class="pie-expand-btn" data-expand-pie="pieResponsible" title="Expandir gráfico">⛶</button>
          <div>
            <h3>Responsable del reporte</h3>
            <p>Participación por responsable de reporte.</p>
          </div>
          <div class="pie-figure" id="pieResponsible"><div class="pie-center"><strong id="pieResponsibleValue">0%</strong><span id="pieResponsibleLabel">Sin datos</span></div></div>
          <div class="pie-legend" id="pieResponsibleLegend"></div>
        </section>
      </div>
    </div>
  </div>

  <script>
    const state = {
      meta: null,
      currentSheet: null,
      rows: [],
      columns: [],
      activeFilters: {},
      latestPayload: null,
    };

    const filterContainer = document.getElementById('filters');
    const kpiGrid = document.getElementById('kpiGrid');
    const tableWrap = document.getElementById('tableWrap');
    const statusChart = document.getElementById('statusChart');
    const countryChart = document.getElementById('countryChart');
    const summaryPills = document.getElementById('summaryPills');
    const loadingState = document.getElementById('loadingState');
    const optionStatus = document.getElementById('optionStatus');
    const tableSubtitle = document.getElementById('tableSubtitle');

    const PIE_COLORS = ['#e4002b', '#111111', '#ff8aa2', '#d0b28f', '#7f1d1d', '#f3a6b5', '#6b7280'];
    const PRIORITY_COLORS = {
      'legal y mandatorio': '#d64545',
      'auditoria': '#b16a00',
      'soporte': '#2e7d32',
      'regulatorio': '#2e7d32',
      'operacional': '#1976d2',
      'mejora': '#7b1fa2',
    };

    const STATUS_COLORS = {
      'produccion': '#2e7d32',
      'production': '#2e7d32',
      'escalado a desarrollo': '#1976d2',
      'escalado a despliegue': '#f57c00',
      'qa': '#ff9800',
      'done': '#2e7d32',
      'solventado': '#2e7d32',
      'cerrado': '#2e7d32',
      'cerrada': '#2e7d32',
      'completado': '#2e7d32',
      'finalizado': '#2e7d32',
      'closed': '#2e7d32',
      'resuelto': '#2e7d32',
      'requerimiento': '#9c27b0',
      'pendiente': '#f44336',
      'en progreso': '#2196f3',
    };

    function apiUrl(path, params = {}) {
      const url = new URL(path, window.location.origin);
      Object.entries(params).forEach(([key, value]) => {
        if (Array.isArray(value)) {
          value.forEach((item) => {
            if (item !== undefined && item !== null && String(item).trim() !== '' && String(item) !== 'all') {
              url.searchParams.append(key, item);
            }
          });
          return;
        }
        if (value !== undefined && value !== null && String(value).trim() !== '' && String(value) !== 'all') {
          url.searchParams.set(key, value);
        }
      });
      return url.toString();
    }

    function escapeHtml(value) {
      return String(value ?? '')
        .replaceAll('&', '&amp;')
        .replaceAll('<', '&lt;')
        .replaceAll('>', '&gt;')
        .replaceAll('"', '&quot;')
        .replaceAll("'", '&#39;');
    }

    function formatValue(value) {
      if (value === null || value === undefined || value === '') return '—';
      return escapeHtml(value);
    }

    function getColumnName(columns, candidates) {
      const lookup = new Map(columns.map((column) => [String(column).toLowerCase(), column]));
      for (const candidate of candidates) {
        const match = lookup.get(String(candidate).toLowerCase());
        if (match) return match;
      }
      return '';
    }

    function countPieData(rows, columnName, limit = 5) {
      if (!columnName) return [];
      const counter = new Map();
      let total = 0;

      rows.forEach((row) => {
        const value = row[columnName];
        if (value === null || value === undefined || String(value).trim() === '') return;
        const key = String(value).trim();
        counter.set(key, (counter.get(key) || 0) + 1);
        total += 1;
      });

      if (!total) return [];

      const sorted = [...counter.entries()].sort((a, b) => b[1] - a[1]);
      const top = sorted.slice(0, limit - 1);
      const items = top.map(([label, count], index) => ({
        label,
        count,
        share: Math.round((count / total) * 1000) / 10,
        color: PIE_COLORS[index % PIE_COLORS.length],
      }));

      const shown = items.reduce((sum, item) => sum + item.count, 0);
      const remaining = total - shown;
      if (remaining > 0) {
        items.push({
          label: 'Otros',
          count: remaining,
          share: Math.round((remaining / total) * 1000) / 10,
          color: '#9ca3af',
        });
      }

      return items;
    }

    function buildPieGradient(items) {
      if (!items.length) return 'conic-gradient(#ddd 0deg 360deg)';
      let start = 0;
      const parts = items.map((item) => {
        const end = start + (item.share / 100) * 360;
        const piece = `${item.color} ${start}deg ${end}deg`;
        start = end;
        return piece;
      });
      return `conic-gradient(${parts.join(', ')})`;
    }

    function renderPieCard(figureId, valueId, labelId, legendId, items, emptyLabel) {
      const figure = document.getElementById(figureId);
      const value = document.getElementById(valueId);
      const label = document.getElementById(labelId);
      const legend = document.getElementById(legendId);

      if (!items.length) {
        figure.style.background = 'conic-gradient(#ddd 0deg 360deg)';
        value.textContent = '0%';
        label.textContent = emptyLabel;
        legend.innerHTML = '<div class="empty">Sin datos para este pastel.</div>';
        return;
      }

      figure.style.background = buildPieGradient(items);
      value.textContent = `${items[0].share}%`;
      label.textContent = `${items[0].label} lidera`;
      legend.innerHTML = items.map((item) => `
        <div class="legend-item">
          <span class="legend-swatch" style="background:${item.color};"></span>
          <span>${escapeHtml(item.label)}</span>
          <span class="legend-meta">${item.count} · ${item.share}%</span>
        </div>
      `).join('');
    }

    function buildPieCardHtml(title, subtitle, items, emptyLabel) {
      if (!items.length) {
        return `
          <section class="pie-card">
            <div><h3>${escapeHtml(title)}</h3><p>${escapeHtml(subtitle)}</p></div>
            <div class="pie-figure" style="background:conic-gradient(#ddd 0deg 360deg);">
              <div class="pie-center"><strong>0%</strong><span>${escapeHtml(emptyLabel)}</span></div>
            </div>
            <div class="pie-legend"><div class="empty">Sin datos para este pastel.</div></div>
          </section>
        `;
      }
      const legend = items.map((item) => `
        <div class="legend-item">
          <span class="legend-swatch" style="background:${item.color};"></span>
          <span>${escapeHtml(item.label)}</span>
          <span class="legend-meta">${item.count} · ${item.share}%</span>
        </div>
      `).join('');
      return `
        <section class="pie-card">
          <div><h3>${escapeHtml(title)}</h3><p>${escapeHtml(subtitle)}</p></div>
          <div class="pie-figure" style="background:${buildPieGradient(items)};">
            <div class="pie-center"><strong>${items[0].share}%</strong><span>${escapeHtml(items[0].label)} lidera</span></div>
          </div>
          <div class="pie-legend">${legend}</div>
        </section>
      `;
    }

    function renderPrioBySprint(payload, sprintColumn, priorityColumn) {
      const grid = document.getElementById('prioBySprintGrid');
      const hint = document.getElementById('prioBySprintHint');
      if (!grid) return;

      const selected = sprintColumn && Array.isArray(state.activeFilters[sprintColumn])
        ? state.activeFilters[sprintColumn].filter((value) => value && value !== 'all')
        : [];

      if (!priorityColumn) {
        grid.innerHTML = '<div class="empty">La hoja actual no tiene el campo prioridad DB.</div>';
        if (hint) hint.textContent = 'Sin campo prioridad DB en esta hoja.';
        return;
      }
      if (!selected.length) {
        grid.innerHTML = '';
        if (hint) hint.textContent = 'Marca uno o más sprints en el filtro para comparar su prioridad DB.';
        return;
      }

      if (hint) hint.textContent = 'Comparando prioridad DB en ' + selected.length + ' sprint(s): ' + selected.join(', ');

      grid.innerHTML = selected.map((sprint, sprintIndex) => {
        const sprintRows = payload.rows.filter((row) => String(row[sprintColumn] ?? '').trim() === String(sprint).trim());
        const items = countPieData(sprintRows, priorityColumn);
        const pieId = `pioPriorityBySprint_${sprintIndex}`;

        const statusColumn = getColumnName(payload.columns, ['Status', 'Estado', 'estado', 'estado de reporte']);
        const tableRows = items.map((item) => {
          const itemRows = sprintRows.filter((row) => String(row[priorityColumn] ?? '').trim() === String(item.label).trim());
          return itemRows.map((row) => {
            const statusValue = String(row[statusColumn] ?? '');
            const statusColor = STATUS_COLORS[statusValue.toLowerCase()] || '#9e9e9e';
            const statusTextColor = ['#2e7d32', '#1976d2', '#7b1fa2'].includes(statusColor) ? '#fff' : '#000';
            return `
              <tr>
                <td>${escapeHtml(String(row['PAIS'] ?? ''))}</td>
                <td>${escapeHtml(String(row['NOVEDAD'] ?? ''))}</td>
                <td>${escapeHtml(String(row['sprint Reportes'] ?? ''))}</td>
                <td>${escapeHtml(String(row['tarjeta Devops'] ?? ''))}</td>
                <td>${escapeHtml(String(row['Sprint despliegue'] ?? ''))}</td>
                <td style="background-color: ${statusColor}; color: ${statusTextColor}; font-weight: 600; padding: 8px; border-radius: 4px;">${escapeHtml(statusValue)}</td>
                <td style="background-color: ${PRIORITY_COLORS[String(row[priorityColumn] ?? '').toLowerCase()] || '#f0f0f0'}; color: #fff; font-weight: 600; text-align: center;">${escapeHtml(String(row[priorityColumn] ?? ''))}</td>
              </tr>
            `;
          }).join('');
        }).join('');

        const tableHtml = sprintRows && sprintRows.length > 0 ? `
          <div style="margin-top: 16px;">
            <h4 style="margin: 0 0 8px; font-size: 13px; color: #666;">Detalle de registros (${sprintRows.length} registros)</h4>
            <div style="border: 1px solid #ddd; border-radius: 8px; overflow-x: auto;">
              <table style="width: 100%; border-collapse: collapse; font-size: 12px;">
                <thead style="background: #f5f5f5;">
                  <tr>
                    <th style="padding: 8px; text-align: left; border-bottom: 1px solid #ddd; font-weight: 600;">País</th>
                    <th style="padding: 8px; text-align: left; border-bottom: 1px solid #ddd; font-weight: 600;">Novedad</th>
                    <th style="padding: 8px; text-align: left; border-bottom: 1px solid #ddd; font-weight: 600;">Sprint Reportes</th>
                    <th style="padding: 8px; text-align: left; border-bottom: 1px solid #ddd; font-weight: 600;">Tarjeta Devops</th>
                    <th style="padding: 8px; text-align: left; border-bottom: 1px solid #ddd; font-weight: 600;">Sprint Despliegue</th>
                    <th style="padding: 8px; text-align: center; border-bottom: 1px solid #ddd; font-weight: 600;">Status</th>
                    <th style="padding: 8px; text-align: center; border-bottom: 1px solid #ddd; font-weight: 600;">Prioridad DB</th>
                  </tr>
                </thead>
                <tbody>
                  ${tableRows}
                </tbody>
              </table>
            </div>
          </div>
        ` : '';

        return `
          <section class="pie-card">
            <button class="pie-expand-btn" data-expand-pie="${pieId}" title="Expandir gráfico">⛶</button>
            <div>
              <h3>Sprint ${sprint}</h3>
              <p>${sprintRows.length} registros</p>
            </div>
            <div class="pie-figure" id="${pieId}" style="background:${buildPieGradient(items)};">
              <div class="pie-center"><strong>${items[0]?.share || 0}%</strong><span>${items[0]?.label || 'Sin datos'}</span></div>
            </div>
            <div class="pie-legend">
              ${items.map((item) => `
                <div class="legend-item">
                  <span class="legend-swatch" style="background:${item.color};"></span>
                  <span>${escapeHtml(item.label)}</span>
                  <span class="legend-meta">${item.count} · ${item.share}%</span>
                </div>
              `).join('')}
            </div>
            ${tableHtml}
          </section>
        `;
      }).join('');
    }

    function renderPastelesSprintSection(payload) {
      const statusColumn = getColumnName(payload.columns, ['Status', 'Estado', 'estado', 'estado de reporte']);
      const countryColumn = getColumnName(payload.columns, ['PAIS', 'País', 'pais', 'country']);
      const sprintColumn = getColumnName(payload.columns, ['Sprint despliegue']);
      const priorityColumn = getColumnName(payload.columns, ['prioridad DB', 'prioridad db']);
      const respColumn = getColumnName(payload.columns, ['Responsable de Reporte', 'Responsable reporte']);

      renderPieCard('pieSprint1', 'pieSprint1Value', 'pieSprint1Label', 'pieSprint1Legend', countPieData(payload.rows, statusColumn), 'Sin estado');
      renderPieCard('pieSprint2', 'pieSprint2Value', 'pieSprint2Label', 'pieSprint2Legend', countPieData(payload.rows, countryColumn), 'Sin país');
      renderPieCard('pieSprint3', 'pieSprint3Value', 'pieSprint3Label', 'pieSprint3Legend', countPieData(payload.rows, sprintColumn), 'Sin sprint despliegue');
      renderPieCard('pieSprint4', 'pieSprint4Value', 'pieSprint4Label', 'pieSprint4Legend', countPieData(payload.rows, priorityColumn), 'Sin prioridad DB');
      renderPieCard('pieSprint5', 'pieSprint5Value', 'pieSprint5Label', 'pieSprint5Legend', countPieData(payload.rows, respColumn), 'Sin responsable');
    }

    function renderInsightModal(payload) {
      const statusColumn = getColumnName(payload.columns, ['Status', 'Estado', 'estado', 'estado de reporte']);
      const countryColumn = getColumnName(payload.columns, ['PAIS', 'País', 'pais', 'country']);
      const sprintColumn = getColumnName(payload.columns, ['Sprint despliegue']);
      const priorityColumn = getColumnName(payload.columns, ['prioridad DB', 'prioridad db']);
      const respColumn = getColumnName(payload.columns, ['Responsable de Reporte', 'Responsable reporte']);

      const sprintLabel = document.getElementById('insightSprintLabel');
      if (sprintLabel) {
        const selected = sprintColumn ? state.activeFilters[sprintColumn] : null;
        const sprints = Array.isArray(selected) ? selected.filter((v) => v && v !== 'all') : [];
        const text = sprints.length ? sprints.join(', ') : 'Todos';
        sprintLabel.innerHTML = `<strong>Sprint despliegue:</strong> ${escapeHtml(text)}`;
      }

      renderPieCard('pieStatus', 'pieStatusValue', 'pieStatusLabel', 'pieStatusLegend', countPieData(payload.rows, statusColumn), 'Sin estado');
      renderPieCard('pieCountry', 'pieCountryValue', 'pieCountryLabel', 'pieCountryLegend', countPieData(payload.rows, countryColumn), 'Sin país');
      renderPieCard('pieDeployment', 'pieDeploymentValue', 'pieDeploymentLabel', 'pieDeploymentLegend', countPieData(payload.rows, sprintColumn), 'Sin sprint despliegue');
      renderPieCard('piePriorityDb', 'piePriorityDbValue', 'piePriorityDbLabel', 'piePriorityDbLegend', countPieData(payload.rows, priorityColumn), 'Sin prioridad DB');
      renderPieCard('pieResponsible', 'pieResponsibleValue', 'pieResponsibleLabel', 'pieResponsibleLegend', countPieData(payload.rows, respColumn), 'Sin responsable');
    }

    function buildCurrentParams() {
      const params = { sheet: state.currentSheet || '' };
      Object.entries(state.activeFilters).forEach(([key, value]) => {
        if (Array.isArray(value) && value.length > 0) {
          params[key] = value;
          return;
        }
        if (value !== undefined && value !== null && String(value).trim() !== '' && String(value) !== 'all') {
          params[key] = value;
        }
      });
      return params;
    }

    function openInsights() {
      if (!state.latestPayload) return;
      renderInsightModal(state.latestPayload);
      const modal = document.getElementById('insightModal');
      modal.classList.add('open');
      modal.setAttribute('aria-hidden', 'false');
    }

    function closeInsights() {
      const modal = document.getElementById('insightModal');
      modal.classList.remove('open');
      modal.setAttribute('aria-hidden', 'true');
    }

    function openPieExpand(pieId) {
      const pieFigure = document.getElementById(pieId);
      if (!pieFigure) return;

      const pieLegendId = pieId + 'Legend';
      const pieLegend = document.getElementById(pieLegendId);

      const pieTitle = pieFigure.closest('.pie-card')?.querySelector('h3')?.textContent || 'Gráfico expandido';
      const pieSubtitle = pieFigure.closest('.pie-card')?.querySelector('p')?.textContent || '';

      const expandFigure = document.getElementById('pieExpandFigure');
      const expandLegend = document.getElementById('pieExpandLegend');
      const expandTitle = document.getElementById('pieExpandTitle');
      const expandSubtitle = document.getElementById('pieExpandSubtitle');

      if (expandTitle) expandTitle.textContent = pieTitle;
      if (expandSubtitle) expandSubtitle.textContent = pieSubtitle;

      if (expandFigure) {
        expandFigure.style.background = pieFigure.style.background;
        const centerHtml = pieFigure.querySelector('.pie-center')?.innerHTML || '';
        expandFigure.innerHTML = `<div class="pie-center">${centerHtml}</div>`;
      }

      if (expandLegend && pieLegend) {
        expandLegend.innerHTML = pieLegend.innerHTML;
      }

      const modal = document.getElementById('pieExpandModal');
      if (modal) {
        modal.classList.add('open');
        modal.setAttribute('aria-hidden', 'false');
      }
    }

    function closePieExpand() {
      const modal = document.getElementById('pieExpandModal');
      if (modal) {
        modal.classList.remove('open');
        modal.setAttribute('aria-hidden', 'true');
      }
    }

    function triggerExport(format) {
      window.location.href = apiUrl(`/export/${format}`, buildCurrentParams());
    }

    async function refreshExcelData() {
      optionStatus.textContent = 'Actualizando desde Excel...';
      try {
        const response = await fetch('/api/reload');
        if (!response.ok) {
          throw new Error('No fue posible actualizar desde Excel');
        }
        const data = await response.json();
        optionStatus.textContent = `Actualizado: ${data.refreshed_at} · ${data.sheet_count} hojas`;
        await loadData(true);
      } catch (error) {
        optionStatus.textContent = error.message;
      }
    }

    function renderPills(meta, rows) {
      const sheet = state.currentSheet || '';
      summaryPills.innerHTML = `
        <span class="pill"><strong>${escapeHtml(sheet)}</strong> hoja activa</span>
        <span class="pill"><strong>${meta.columns.length}</strong> columnas</span>
        <span class="pill"><strong>${rows.length}</strong> registros visibles</span>
        <span class="pill"><strong>${meta.sheets.length}</strong> hojas del libro</span>
      `;
    }

    function renderKpis(meta) {
      const items = meta.kpis || [];
      kpiGrid.innerHTML = items.map((item) => `
        <article class="panel kpi">
          <div class="label">${escapeHtml(item.label)}</div>
          <div class="value">${escapeHtml(item.value)}</div>
          <div class="hint">${escapeHtml(item.hint || '')}</div>
        </article>
      `).join('');
    }

    function renderBars(target, items) {
      if (!items.length) {
        target.innerHTML = '<div class="empty">No hay datos para mostrar con los filtros actuales.</div>';
        return;
      }
      const max = Math.max(...items.map((item) => item.count), 1);
      target.innerHTML = items.map((item) => {
        const width = Math.max(5, Math.round((item.count / max) * 100));
        return `
          <div class="bar-item">
            <div class="bar-top">
              <span>${escapeHtml(item.label)}</span>
              <span><strong>${item.count}</strong></span>
            </div>
            <div class="bar-track"><div class="bar-fill" style="width:${width}%;"></div></div>
          </div>
        `;
      }).join('');
    }

    function renderTable(columns, rows) {
      if (!columns.length) {
        tableWrap.innerHTML = '<div class="empty">La hoja seleccionada no tiene columnas válidas.</div>';
        return;
      }
      if (!rows.length) {
        tableWrap.innerHTML = '<div class="empty">No hay registros con esos filtros.</div>';
        return;
      }

      const hidden = [
        'version', 'fecha reporte', 'fecha despliegue',
        '#', 'despliegue tentativo', 'caso aranda', 'reduccion', 'tipo', 'observacion', 'prioridad',
      ];
      const visibleColumns = columns
        .filter((col) => !hidden.includes(String(col).trim().toLowerCase()))
        .slice(0, 18);
      const head = `<th style="width:50px;text-align:center;">Seq</th>` + visibleColumns.map((col) => `<th>${escapeHtml(col)}</th>`).join('');

      const priorityColName = getColumnName(columns, ['prioridad db', 'prioridad DB']);
      const body = rows.map((row, index) => {
        const seqCell = `<td style="text-align:center;font-weight:600;color:var(--muted);">${index + 1}</td>`;
        const cells = visibleColumns.map((col) => {
          const value = formatValue(row[col]);
          if (priorityColName && col.toLowerCase() === priorityColName.toLowerCase()) {
            const priorityKey = String(row[col] ?? '').trim().toLowerCase();
            const color = PRIORITY_COLORS[priorityKey] || 'transparent';
            const textColor = ['#e4002b', '#b16a00'].includes(color) ? '#fff' : '#000';
            return `<td style="background-color:${color};color:${textColor};font-weight:600;padding:10px;border-radius:4px;">${value}</td>`;
          }
          return `<td>${value}</td>`;
        }).join('');
        return `<tr>${seqCell}${cells}</tr>`;
      }).join('');

      tableWrap.innerHTML = `
        <table>
          <thead><tr>${head}</tr></thead>
          <tbody>${body}</tbody>
        </table>
      `;
    }

    function renderFilters(meta) {
      const sheetOptions = meta.sheets
        .map((sheet) => `<option value="${escapeHtml(sheet)}" ${sheet === state.currentSheet ? 'selected' : ''}>${escapeHtml(sheet)}</option>`)
        .join('');

      const filterOptions = meta.availableFilters
        .map((filter) => {
          const currentValue = state.activeFilters[filter.name];
          const isSprintMulti = String(filter.name).toLowerCase() === 'sprint despliegue';
          const selectedValues = Array.isArray(currentValue) ? currentValue.map(String) : [String(currentValue || 'all')];

          if (isSprintMulti) {
            const marked = selectedValues.filter((value) => value && value !== 'all');
            const summary = marked.length ? `${marked.length} sprint(s) marcado(s)` : 'Todos los sprints';
            const options = filter.values.map((value) => {
              const checked = marked.includes(String(value)) ? 'checked' : '';
              return `
                <label class="combo-option">
                  <input type="checkbox" value="${escapeHtml(value)}" ${checked} />
                  <span>${escapeHtml(value)}</span>
                </label>
              `;
            }).join('');
            return `
              <div class="field">
                <label>${escapeHtml(filter.label)}</label>
                <div class="combo" data-combo="${escapeHtml(filter.name)}">
                  <button type="button" class="combo-toggle" data-combo-toggle>
                    <span data-combo-summary>${summary}</span>
                    <span class="chev">▼</span>
                  </button>
                  <div class="combo-panel">${options}</div>
                </div>
                <span class="hint">Marca los sprints que necesitas visualizar.</span>
              </div>
            `;
          }

          return `
            <div class="field">
              <label for="filter-${escapeHtml(filter.name)}">${escapeHtml(filter.label)}</label>
              <select id="filter-${escapeHtml(filter.name)}" data-filter="${escapeHtml(filter.name)}">
                <option value="all">Todos</option>
                ${filter.values.map((value) => `<option value="${escapeHtml(value)}" ${String(currentValue || 'all') === String(value) ? 'selected' : ''}>${escapeHtml(value)}</option>`).join('')}
              </select>
            </div>
          `;
        })
        .join('');

      filterContainer.innerHTML = `
        <div class="field">
          <label for="sheetSelect">Hoja</label>
          <select id="sheetSelect">${sheetOptions}</select>
        </div>
        ${filterOptions}
        <div style="display: grid; grid-template-columns: 1fr auto; gap: 10px; align-items: flex-end;">
          <div class="field">
            <label for="searchInput">Búsqueda</label>
            <input id="searchInput" type="search" placeholder="Buscar en la hoja seleccionada" value="${escapeHtml(state.activeFilters.q || '')}" />
          </div>
          <button class="option-btn secondary" id="clearFiltersBtn" type="button" style="margin-bottom: 0;">Limpiar filtros</button>
        </div>
      `;

      document.getElementById('sheetSelect').addEventListener('change', async (event) => {
        state.currentSheet = event.target.value;
        state.activeFilters = { q: '' };
        await loadData();
      });

      document.getElementById('searchInput').addEventListener('input', debounce(async (event) => {
        state.activeFilters.q = event.target.value;
        await loadData(false);
      }, 220));

      filterContainer.querySelectorAll('[data-combo]').forEach((combo) => {
        const name = combo.getAttribute('data-combo');
        const toggle = combo.querySelector('[data-combo-toggle]');
        const summaryEl = combo.querySelector('[data-combo-summary]');
        const panel = combo.querySelector('.combo-panel');
        const checkboxes = () => Array.from(panel.querySelectorAll('input[type="checkbox"]'));
        const selected = () => checkboxes().filter((cb) => cb.checked).map((cb) => cb.value);

        toggle.addEventListener('click', (event) => {
          event.stopPropagation();
          const isOpen = combo.classList.contains('open');
          document.querySelectorAll('.combo.open').forEach((c) => c.classList.remove('open'));
          if (isOpen) {
            loadData(false);
          } else {
            combo.classList.add('open');
          }
        });

        panel.addEventListener('click', (event) => event.stopPropagation());

        checkboxes().forEach((cb) => {
          cb.addEventListener('change', () => {
            const marked = selected();
            state.activeFilters[name] = marked;
            if (summaryEl) {
              summaryEl.textContent = marked.length ? `${marked.length} sprint(s) marcado(s)` : 'Todos los sprints';
            }
          });
        });
      });

      meta.availableFilters.forEach((filter) => {
        const select = document.getElementById(`filter-${filter.name}`);
        if (!select) return;
        select.addEventListener('change', async (event) => {
          state.activeFilters[filter.name] = event.target.value;
          await loadData(false);
        });
      });

      const clearBtn = document.getElementById('clearFiltersBtn');
      if (clearBtn) {
        clearBtn.addEventListener('click', async () => {
          state.currentSheet = state.meta?.default_sheet || state.currentSheet;
          state.activeFilters = { q: '' };
          await loadData(true);
        });
      }
    }

    function debounce(fn, delay) {
      let timer = null;
      return (...args) => {
        window.clearTimeout(timer);
        timer = window.setTimeout(() => fn(...args), delay);
      };
    }

    async function loadMeta(initial = false) {
      const response = await fetch(apiUrl('/api/meta', { sheet: state.currentSheet || '' }));
      if (!response.ok) {
        throw new Error('No se pudo cargar metadata');
      }
      const meta = await response.json();
      state.meta = meta;
      if (initial || !state.currentSheet) {
        state.currentSheet = meta.default_sheet;
      }
      return meta;
    }

    async function loadData(resetMeta = true) {
      loadingState.style.display = 'inline-flex';
      try {
        const isInitialLoad = resetMeta && !state.meta;
        if (resetMeta || !state.meta) {
          await loadMeta();
        }

        if (isInitialLoad && !state.activeFilters.anio) {
          const currentYear = new Date().getFullYear().toString();
          state.activeFilters.anio = currentYear;
        }

        const response = await fetch(apiUrl('/api/data', buildCurrentParams()));
        if (!response.ok) {
          throw new Error('No se pudo cargar data');
        }
        const payload = await response.json();
        state.rows = payload.rows;
        state.columns = payload.columns;
        state.latestPayload = payload;

        renderFilters(payload.meta);
        renderKpis(payload.meta);
        renderPills(payload.meta, payload.rows);
        renderBars(statusChart, payload.status_breakdown);
        renderBars(countryChart, payload.country_breakdown);
        renderTable(payload.columns, payload.rows);
        renderInsightModal(payload);

        const sprintColumn = getColumnName(payload.columns, ['Sprint despliegue']);
        const priorityColumn = getColumnName(payload.columns, ['prioridad DB', 'prioridad db']);
        renderPrioBySprint(payload, sprintColumn, priorityColumn);

        tableSubtitle.textContent = `${payload.rows.length} registros visibles en ${payload.sheet}`;
        optionStatus.textContent = `Sincronizado · hoja ${payload.sheet} · ${payload.rows.length} registros`;
      } catch (error) {
        tableWrap.innerHTML = `<div class="empty">${escapeHtml(error.message)}</div>`;
        optionStatus.textContent = error.message;
      } finally {
        loadingState.style.display = 'none';
      }
    }

    document.addEventListener('click', () => {
      const openCombo = document.querySelector('.combo.open');
      if (openCombo) {
        openCombo.classList.remove('open');
        loadData(false);
      }
    });

    const openInsBtn = document.getElementById('openInsightsBtn');
    if (openInsBtn) openInsBtn.addEventListener('click', openInsights);

    const closeInsBtn = document.getElementById('closeInsightsBtn');
    if (closeInsBtn) closeInsBtn.addEventListener('click', closeInsights);

    const closePieExpandBtn = document.getElementById('closePieExpandBtn');
    if (closePieExpandBtn) closePieExpandBtn.addEventListener('click', closePieExpand);

    document.addEventListener('click', (event) => {
      const btn = event.target.closest('.pie-expand-btn');
      if (btn) {
        event.stopPropagation();
        const pieId = btn.getAttribute('data-expand-pie');
        if (pieId) openPieExpand(pieId);
      }
    });

    const pieExpandModal = document.getElementById('pieExpandModal');
    if (pieExpandModal) {
      pieExpandModal.addEventListener('click', (event) => {
        if (event.target.id === 'pieExpandModal') closePieExpand();
      });
    }

    document.addEventListener('keydown', (event) => {
      if (event.key === 'Escape') {
        closeInsights();
        closePieExpand();
      }
    });

    const savePdfBtn = document.getElementById('savePiesPdfBtn');
    if (savePdfBtn) savePdfBtn.addEventListener('click', () => window.print());

    const pasteleSprintBtn = document.getElementById('openPastelesSprint');
    if (pasteleSprintBtn) {
      pasteleSprintBtn.addEventListener('click', () => {
        const section = document.getElementById('pasteleSprintSection');
        const other = document.getElementById('pasteleProitySection');
        if (other) other.classList.add('hidden-section');
        if (section) section.classList.remove('hidden-section');
        if (section) section.scrollIntoView({ behavior: 'smooth' });
        if (state.latestPayload) {
          renderPastelesSprintSection(state.latestPayload);
        }
      });
    }

    const pasteleProityBtn = document.getElementById('openPastelesPrority');
    if (pasteleProityBtn) {
      pasteleProityBtn.addEventListener('click', () => {
        const section = document.getElementById('pasteleProitySection');
        const other = document.getElementById('pasteleSprintSection');
        if (other) other.classList.add('hidden-section');
        if (section) section.classList.remove('hidden-section');
        if (section) section.scrollIntoView({ behavior: 'smooth' });
      });
    }

    const navTableroGeneral = document.getElementById('navTableroGeneral');
    if (navTableroGeneral) {
      navTableroGeneral.addEventListener('click', () => {
        const pageTablero = document.getElementById('pageTableroGeneral');
        const pagePastelesSprint = document.getElementById('pasteleSprintSection');
        const pagePastelesProity = document.getElementById('pasteleProitySection');
        const pastelesSubmenu = document.getElementById('pastelesSubmenu');
        const pastelesChev = document.getElementById('pastelesChev');

        if (pageTablero) pageTablero.classList.remove('hidden-section');
        if (pagePastelesSprint) pagePastelesSprint.classList.add('hidden-section');
        if (pagePastelesProity) pagePastelesProity.classList.add('hidden-section');
        if (pastelesSubmenu) pastelesSubmenu.classList.add('hidden');
        if (pastelesChev) pastelesChev.textContent = '▶';

        window.scrollTo({ top: 0, behavior: 'smooth' });
      });
    }


    const navPastelesToggle = document.getElementById('navPastelesToggle');
    if (navPastelesToggle) {
      navPastelesToggle.addEventListener('click', () => {
        const pastelesSubmenu = document.getElementById('pastelesSubmenu');
        const pastelesChev = document.getElementById('pastelesChev');

        if (pastelesSubmenu) {
          pastelesSubmenu.classList.toggle('hidden');
          if (pastelesChev) {
            const isOpen = !pastelesSubmenu.classList.contains('hidden');
            pastelesChev.textContent = isOpen ? '▼' : '▶';
            pastelesChev.style.transition = 'transform 0.3s ease';
          }
        }
      });
    }

    const navPastelesSprint = document.getElementById('navPastelesSprint');
    if (navPastelesSprint) {
      navPastelesSprint.addEventListener('click', () => {
        const pageTablero = document.getElementById('pageTableroGeneral');
        const pagePastelesSprint = document.getElementById('pasteleSprintSection');
        const pagePastelesProity = document.getElementById('pasteleProitySection');

        if (pageTablero) pageTablero.classList.add('hidden-section');
        if (pagePastelesSprint) pagePastelesSprint.classList.remove('hidden-section');
        if (pagePastelesProity) pagePastelesProity.classList.add('hidden-section');

        if (pagePastelesSprint) {
          pagePastelesSprint.scrollIntoView({ behavior: 'smooth' });
          if (state.latestPayload) {
            renderPastelesSprintSection(state.latestPayload);
          }
        }
      });
    }


    const navPastelesProity = document.getElementById('navPastelesPrority');
    if (navPastelesProity) {
      navPastelesProity.addEventListener('click', () => {
        const pageTablero = document.getElementById('pageTableroGeneral');
        const pagePastelesSprint = document.getElementById('pasteleSprintSection');
        const pagePastelesProity = document.getElementById('pasteleProitySection');

        if (pageTablero) pageTablero.classList.add('hidden-section');
        if (pagePastelesSprint) pagePastelesSprint.classList.add('hidden-section');
        if (pagePastelesProity) pagePastelesProity.classList.remove('hidden-section');

        if (pagePastelesProity) {
          pagePastelesProity.scrollIntoView({ behavior: 'smooth' });
        }
      });
    }


    document.getElementById('exportExcelBtn').addEventListener('click', () => triggerExport('excel'));
    document.getElementById('exportPdfBtn').addEventListener('click', () => triggerExport('pdf'));
    document.getElementById('refreshExcelBtn').addEventListener('click', refreshExcelData);

    const savePrioDbPdfBtn = document.getElementById('savePrioDbPdfBtn');
    if (savePrioDbPdfBtn) {
      savePrioDbPdfBtn.addEventListener('click', () => window.print());
    }

    const savePrioDbExcelBtn = document.getElementById('savePrioDbExcelBtn');
    if (savePrioDbExcelBtn) {
      savePrioDbExcelBtn.addEventListener('click', () => triggerExport('excel'));
    }

    const refreshExcelBtnMain = document.getElementById('refreshExcelBtnMain');
    if (refreshExcelBtnMain) {
      refreshExcelBtnMain.addEventListener('click', refreshExcelData);
    }

    document.getElementById('insightModal').addEventListener('click', (event) => {
      if (event.target.id === 'insightModal') closeInsights();
    });

    loadData(true);
  </script>
</body>
</html>
"""


def convert_value(value: object) -> object:
    if isinstance(value, (datetime, date)):
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M")
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        return round(value, 6)
    return value


def load_workbook_data() -> dict[str, dict[str, object]]:
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"No se encontro el Excel en {EXCEL_PATH}")

    workbook = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    workbook_data: dict[str, dict[str, object]] = {}

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            workbook_data[worksheet.title] = {"columns": [], "rows": []}
            continue

        header_row = rows[0]
        last_non_empty = -1
        for index, value in enumerate(header_row):
            if value is not None and str(value).strip() != "":
                last_non_empty = index

        columns = []
        for index, value in enumerate(header_row[: last_non_empty + 1]):
            if value is None or str(value).strip() == "":
                columns.append(f"Columna {index + 1}")
            else:
                columns.append(str(value).strip())

        records = []
        for raw_row in rows[1:]:
            limited = raw_row[: len(columns)]
            if not any(cell is not None and str(cell).strip() != "" for cell in limited):
                continue
            record = {}
            for index, column_name in enumerate(columns):
                value = limited[index] if index < len(limited) else None
                record[column_name] = convert_value(value)
            records.append(record)

        workbook_data[worksheet.title] = {"columns": columns, "rows": records}

    return workbook_data


def refresh_workbook_data() -> dict[str, object]:
    global WORKBOOK_DATA, DEFAULT_SHEET, LAST_REFRESH
    WORKBOOK_DATA = load_workbook_data()
    DEFAULT_SHEET = "SIR" if "SIR" in WORKBOOK_DATA else next(iter(WORKBOOK_DATA), "")
    LAST_REFRESH = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return {
        "sheet_count": len(WORKBOOK_DATA),
        "default_sheet": DEFAULT_SHEET,
        "refreshed_at": LAST_REFRESH,
    }


def get_sheet_data(sheet_name: str) -> dict[str, object]:
    if sheet_name in WORKBOOK_DATA:
        return WORKBOOK_DATA[sheet_name]
    if DEFAULT_SHEET:
        return WORKBOOK_DATA[DEFAULT_SHEET]
    return {"columns": [], "rows": []}


def matches_filter(record: dict[str, object], column: str, value: str) -> bool:
    if value in {"", "all", None}:
        return True
    record_value = record.get(column)
    if record_value is None:
        return False
    options = [item.strip().lower() for item in str(value).split("||") if item.strip()]
    if not options:
        return True
    return str(record_value).strip().lower() in options


def apply_filters(rows: list[dict[str, object]], query: dict[str, str]) -> list[dict[str, object]]:
    search = query.get("q", "").strip().lower()
    filtered_rows = []

    for row in rows:
        if search:
            haystack = " ".join(str(value).lower() for value in row.values() if value not in {None, ""})
            if search not in haystack:
                continue

        passed = True
        for key, value in query.items():
            if key in {"sheet", "q"} or value in {"", "all", None}:
                continue
            if key not in row:
                continue
            if not matches_filter(row, key, value):
                passed = False
                break
        if passed:
            filtered_rows.append(row)

    return filtered_rows


def pick_column(columns: list[str], candidates: list[str]) -> str | None:
    lowered = {column.lower(): column for column in columns}
    for candidate in candidates:
        if candidate.lower() in lowered:
            return lowered[candidate.lower()]
    return None


def build_breakdown(rows: list[dict[str, object]], column_name: str | None, limit: int = 8) -> list[dict[str, object]]:
    if not column_name:
        return []

    counter = Counter()
    for row in rows:
        value = row.get(column_name)
        if value in {None, ""}:
            continue
        counter[str(value)] += 1

    total = sum(counter.values()) or 1
    result = []
    for label, count in counter.most_common(limit):
        result.append({
            "label": label,
            "count": count,
            "share": round((count / total) * 100, 1),
        })
    return result


def build_available_filters(columns: list[str], rows: list[dict[str, object]]) -> list[dict[str, object]]:
    available = []
    for candidate in CANDIDATE_FILTERS:
        column_name = pick_column(columns, [candidate])
        if not column_name:
            continue

        seen = set()
        values = []
        for row in rows:
            value = row.get(column_name)
            if value in {None, ""}:
                continue
            value_text = str(value)
            key = value_text.strip().lower()
            if key in seen:
                continue
            seen.add(key)
            values.append(value_text)

        values = sorted(values, key=lambda item: item.lower())[:50]
        available.append({"name": column_name, "label": column_name, "values": values})
    return available


def status_summary(rows: list[dict[str, object]], columns: list[str]) -> tuple[int, int, int, str | None]:
    status_column = pick_column(columns, ["Status", "estado", "estado de reporte"])
    if not status_column:
        return len(rows), len(rows), 0, None

    open_count = 0
    closed_count = 0
    for row in rows:
        value = row.get(status_column)
        if value in {None, ""}:
            open_count += 1
            continue
        if str(value).strip().lower() in CLOSED_STATUS_VALUES:
            closed_count += 1
        else:
            open_count += 1

    return len(rows), open_count, closed_count, status_column


def unique_count(rows: list[dict[str, object]], column_name: str | None) -> int:
    if not column_name:
        return 0
    values = {str(row.get(column_name)).strip().lower() for row in rows if row.get(column_name) not in {None, ""}}
    return len(values)


def top_label(rows: list[dict[str, object]], column_name: str | None) -> tuple[str, int]:
    if not column_name:
        return "N/A", 0
    counter = Counter()
    for row in rows:
        value = row.get(column_name)
        if value in {None, ""}:
            continue
        counter[str(value)] += 1
    if not counter:
        return "N/A", 0
    label, count = counter.most_common(1)[0]
    return label, count


def build_meta(sheet_name: str, rows: list[dict[str, object]], columns: list[str]) -> dict[str, object]:
    total_records, open_count, closed_count, status_column = status_summary(rows, columns)
    country_column = pick_column(columns, ["PAIS", "País", "pais", "country"])
    resp_column = pick_column(columns, ["Responsable de Reporte", "Responsable reporte"])
    sprint_column = pick_column(columns, ["Sprint despliegue"])
    priority_db_column = pick_column(columns, ["prioridad DB", "prioridad db"])

    top_country, top_country_count = top_label(rows, country_column)
    top_resp, top_resp_count = top_label(rows, resp_column)
    top_priority, top_priority_count = top_label(rows, priority_db_column)

    return {
        "sheet": sheet_name,
        "sheets": list(WORKBOOK_DATA.keys()),
        "default_sheet": DEFAULT_SHEET,
        "refreshed_at": LAST_REFRESH,
        "columns": columns,
        "total_rows": total_records,
        "kpis": [
            {
                "label": "Registros visibles",
                "value": str(total_records),
                "hint": f"Hoja {sheet_name}",
            },
            {
                "label": "Defectos abiertos",
                "value": str(open_count),
                "hint": f"Segun {status_column or 'status'}",
            },
            {
                "label": "Defectos cerrados",
                "value": str(closed_count),
                "hint": "Conteo en vista filtrada",
            },
            {
                "label": "Paises impactados",
                "value": str(unique_count(rows, country_column)),
                "hint": f"Principal: {top_country} ({top_country_count})",
            },
            {
                "label": "Responsables activos",
                "value": str(unique_count(rows, resp_column)),
                "hint": f"Top: {top_resp} ({top_resp_count})",
            },
            {
                "label": "Sprints despliegue",
                "value": str(unique_count(rows, sprint_column)),
                "hint": f"Campo: {sprint_column or 'N/A'}",
            },
            {
                "label": "Prioridad DB lider",
                "value": top_priority,
                "hint": f"{top_priority_count} registros",
            },
        ],
        "availableFilters": build_available_filters(columns, rows),
    }


def parse_query_params(raw_query: dict[str, list[str]]) -> dict[str, str]:
    params: dict[str, str] = {}
    for key, values in raw_query.items():
        if not values:
            params[key] = ""
        elif len(values) == 1:
            params[key] = values[0]
        else:
            params[key] = "||".join(values)
    return params


def build_context(query: dict[str, str]) -> dict[str, object]:
    sheet_name = query.get("sheet") or DEFAULT_SHEET
    sheet_data = get_sheet_data(sheet_name)
    all_rows = sheet_data["rows"]
    rows = apply_filters(all_rows, query)
    columns = sheet_data["columns"]

    country_column = pick_column(columns, ["PAIS", "País", "pais", "country"])
    status_column = pick_column(columns, ["Status", "estado", "estado de reporte"])
    sprint_column = pick_column(columns, ["Sprint despliegue"])
    priority_db_column = pick_column(columns, ["prioridad DB", "prioridad db"])
    responsible_column = pick_column(columns, ["Responsable de Reporte", "Responsable reporte"])

    meta = build_meta(sheet_name, rows, columns)

    return {
        "sheet": sheet_name,
        "rows": rows,
        "columns": columns,
        "meta": meta,
        "status_breakdown": build_breakdown(rows, status_column),
        "country_breakdown": build_breakdown(rows, country_column),
        "sprint_deployment_breakdown": build_breakdown(rows, sprint_column),
        "priority_db_breakdown": build_breakdown(rows, priority_db_column),
        "responsible_breakdown": build_breakdown(rows, responsible_column),
    }


def safe_filename_part(value: str) -> str:
    safe = []
    for char in value:
        if char.isalnum() or char in {"-", "_"}:
            safe.append(char)
        else:
            safe.append("_")
    return "".join(safe).strip("_") or "export"


def csv_safe_text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def build_excel_export(context: dict[str, object]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"

    sheet_name = context["sheet"]
    rows = context["rows"]
    columns = context["columns"]
    meta = context["meta"]

    summary.append(["Tablero gerencial de defectos"])
    summary.append(["Hoja", sheet_name])
    summary.append(["Actualizado", meta["refreshed_at"]])
    summary.append(["Registros visibles", len(rows)])
    summary.append(["Columnas", len(columns)])

    summary.append([])
    summary.append(["KPIs"])
    summary.append(["Indicador", "Valor", "Detalle"])
    for item in meta["kpis"]:
        summary.append([item["label"], item["value"], item["hint"]])

    for section_name, breakdown in [
        ("Defecto por pais", context["country_breakdown"]),
        ("Despliegue por sprint", context["sprint_deployment_breakdown"]),
        ("Prioridad DB", context["priority_db_breakdown"]),
        ("Responsable del reporte", context["responsible_breakdown"]),
    ]:
        summary.append([])
        summary.append([section_name])
        summary.append(["Valor", "Cuenta", "%"])
        for item in breakdown:
            summary.append([item["label"], item["count"], item["share"]])

    data_sheet = workbook.create_sheet("Datos")
    data_sheet.append(columns)
    for row in rows:
        data_sheet.append([csv_safe_text(row.get(col)) for col in columns])

    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 32
    summary.column_dimensions["C"].width = 60

    for col_cells in data_sheet.columns:
        values = [csv_safe_text(cell.value) for cell in col_cells]
        width = min(max((len(v) for v in values), default=8) + 2, 42)
        data_sheet.column_dimensions[col_cells[0].column_letter].width = width

    sprint_column = pick_column(columns, ["Sprint despliegue"])
    priority_column = pick_column(columns, ["prioridad DB", "prioridad db"])
    country_column = pick_column(columns, ["PAIS", "País", "pais", "country"])
    detail_country_column = pick_column(columns, ["Detalle paises"])
    sprint_report_column = pick_column(columns, ["sprint Reportes"])
    card_devops_column = pick_column(columns, ["tarjeta Devops"])

    if sprint_column and priority_column:
        priority_sheet = workbook.create_sheet("Análisis Prioridad DB")

        sprints = sorted(set(str(row.get(sprint_column, "")).strip() for row in rows if row.get(sprint_column)))

        for sprint in sprints:
            sprint_rows = [row for row in rows if str(row.get(sprint_column, "")).strip() == sprint]

            priority_sheet.append([f"Sprint {sprint}"])
            priority_sheet.append(["Resumen"])
            priority_sheet.append(["Prioridad", "Cantidad", "Porcentaje"])

            priority_counts = {}
            for row in sprint_rows:
                priority = str(row.get(priority_column, "Sin prioridad")).strip()
                priority_counts[priority] = priority_counts.get(priority, 0) + 1

            total = sum(priority_counts.values()) or 1
            for priority, count in sorted(priority_counts.items(), key=lambda x: x[1], reverse=True):
                percentage = (count / total * 100) if total > 0 else 0
                priority_sheet.append([priority, count, f"{percentage:.1f}%"])

            priority_sheet.append([])
            priority_sheet.append(["Detalle de Registros"])
            priority_sheet.append(["País", "Novedad", "Sprint Reportes", "Tarjeta Devops", "Sprint Despliegue", "Prioridad DB"])

            novedad_column = pick_column(columns, ["NOVEDAD", "Novedad"])
            for row in sprint_rows:
                priority_sheet.append([
                    csv_safe_text(row.get(country_column, "")),
                    csv_safe_text(row.get(novedad_column, "")),
                    csv_safe_text(row.get(sprint_report_column, "")),
                    csv_safe_text(row.get(card_devops_column, "")),
                    csv_safe_text(row.get(sprint_column, "")),
                    csv_safe_text(row.get(priority_column, "")),
                ])

            priority_sheet.append([])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def pdf_text(value: object) -> str:
    return csv_safe_text(value).encode("latin-1", "replace").decode("latin-1")


def wrap_pdf_line(text: str, width: int = 95) -> list[str]:
    return textwrap.wrap(text, width=width) or [""]


def pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf_bytes(
    pages: list[list[str]],
    page_width: int,
    page_height: int,
    left_margin: int,
    top_margin: int,
    line_height: int,
) -> bytes:
    font_id = 1
    content_base_id = 2
    page_base_id = content_base_id + len(pages)
    pages_id = page_base_id + len(pages)
    catalog_id = pages_id + 1

    objects: dict[int, str] = {
        font_id: "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }

    for index, page in enumerate(pages):
        content_lines = ["BT", "/F1 11 Tf"]
        y = page_height - top_margin - 14
        for line in page:
            safe_line = pdf_escape(line)
            content_lines.append(f"1 0 0 1 {left_margin} {y} Tm ({safe_line}) Tj")
            y -= line_height
        content_lines.append("ET")
        stream = "\n".join(content_lines)
        stream_bytes = stream.encode("latin-1", "replace")
        objects[content_base_id + index] = (
            f"<< /Length {len(stream_bytes)} >>\nstream\n{stream_bytes.decode('latin-1')}\nendstream"
        )

    page_ids = []
    for index in range(len(pages)):
        page_id = page_base_id + index
        page_ids.append(page_id)
        content_id = content_base_id + index
        objects[page_id] = (
            f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 {page_width} {page_height}] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>"
        )

    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects[pages_id] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>"
    objects[catalog_id] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>"

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]

    for obj_id in range(1, catalog_id + 1):
        offsets.append(len(output))
        body = objects[obj_id].encode("latin-1", "replace")
        output.extend(f"{obj_id} 0 obj\n".encode("latin-1"))
        output.extend(body)
        output.extend(b"\nendobj\n")

    xref_offset = len(output)
    output.extend(f"xref\n0 {catalog_id + 1}\n".encode("latin-1"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("latin-1"))

    output.extend(
        (
            f"trailer\n<< /Size {catalog_id + 1} /Root {catalog_id} 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF"
        ).encode("latin-1")
    )

    return bytes(output)


def build_pdf_export(context: dict[str, object]) -> bytes:
    lines: list[str] = []
    meta = context["meta"]
    rows = context["rows"]
    columns = context["columns"]

    lines.append("Tablero gerencial de defectos")
    lines.append(f"Hoja: {context['sheet']}")
    lines.append(f"Actualizado: {meta['refreshed_at']}")
    lines.append(f"Registros visibles: {len(rows)}")
    lines.append("")

    lines.append("KPIs")
    for item in meta["kpis"]:
        lines.append(f"- {item['label']}: {item['value']} ({item['hint']})")

    def add_section(title: str, breakdown: list[dict[str, object]]) -> None:
        lines.append("")
        lines.append(title)
        if not breakdown:
            lines.append("- Sin datos")
            return
        for item in breakdown[:10]:
            lines.append(f"- {item['label']}: {item['count']} ({item['share']}%)")

    add_section("Defecto por pais", context["country_breakdown"])
    add_section("Despliegue por sprint", context["sprint_deployment_breakdown"])
    add_section("Prioridad DB", context["priority_db_breakdown"])
    add_section("Responsable del reporte", context["responsible_breakdown"])

    lines.append("")
    lines.append("Primeros registros")
    preview_cols = columns[:10]
    lines.append(" | ".join(preview_cols))
    for row in rows[:15]:
        lines.append(" | ".join(pdf_text(row.get(col))[:30] for col in preview_cols))

    page_width = 595
    page_height = 842
    left_margin = 40
    top_margin = 40
    line_height = 14
    max_lines = (page_height - top_margin * 2) // line_height

    pages: list[list[str]] = []
    current_page: list[str] = []
    for raw_line in lines:
        for line in wrap_pdf_line(pdf_text(raw_line), width=95):
            if len(current_page) >= max_lines:
                pages.append(current_page)
                current_page = []
            current_page.append(line)
    if current_page:
        pages.append(current_page)

    if not pages:
        pages = [["Sin datos"]]

    return build_pdf_bytes(pages, page_width, page_height, left_margin, top_margin, line_height)


class RequestHandler(BaseHTTPRequestHandler):
    def _send_common_headers(self, content_length: int | None = None) -> None:
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Vary", "Origin")
        if content_length is not None:
            self.send_header("Content-Length", str(content_length))

    def _send_json(self, payload: dict[str, object], status: int = 200) -> None:
        encoded = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self._send_common_headers(len(encoded))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(encoded)

    def _send_html(self, content: str, status: int = 200) -> None:
        encoded = content.encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self._send_common_headers(len(encoded))
        self.end_headers()
        self.wfile.write(encoded)

    def _send_bytes(self, payload: bytes, content_type: str, filename: str) -> None:
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self._send_common_headers(len(payload))
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(payload)

    def do_OPTIONS(self) -> None:  # noqa: N802
        self.send_response(204)
        self._send_common_headers(0)
        self.end_headers()

    def do_GET(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        route = parsed.path
        query = parse_query_params(parse_qs(parsed.query))

        if route == "/":
            self._send_html(HTML_PAGE)
            return

        if route == "/api/meta":
            sheet_name = query.get("sheet") or DEFAULT_SHEET
            sheet_data = get_sheet_data(sheet_name)
            meta = build_meta(sheet_name, sheet_data["rows"], sheet_data["columns"])
            self._send_json(meta)
            return

        if route == "/api/data":
            context = build_context(query)
            payload = {
                "sheet": context["sheet"],
                "columns": context["columns"],
                "rows": context["rows"],
                "meta": context["meta"],
                "status_breakdown": context["status_breakdown"],
                "country_breakdown": context["country_breakdown"],
                "sprint_deployment_breakdown": context["sprint_deployment_breakdown"],
                "priority_db_breakdown": context["priority_db_breakdown"],
                "responsible_breakdown": context["responsible_breakdown"],
            }
            self._send_json(payload)
            return

        if route == "/api/reload":
            refresh_info = refresh_workbook_data()
            self._send_json({"ok": True, **refresh_info})
            return

        if route == "/export/excel":
            context = build_context(query)
            sheet_part = safe_filename_part(str(context["sheet"]))
            payload = build_excel_export(context)
            self._send_bytes(payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"tablero_{sheet_part}.xlsx")
            return

        if route == "/export/pdf":
            context = build_context(query)
            sheet_part = safe_filename_part(str(context["sheet"]))
            payload = build_pdf_export(context)
            self._send_bytes(payload, "application/pdf", f"tablero_{sheet_part}.pdf")
            return

        if route == "/health":
            self._send_json({
                "ok": True,
                "excel_found": EXCEL_PATH.exists(),
                "sheet_count": len(WORKBOOK_DATA),
                "refreshed_at": LAST_REFRESH,
            })
            return

        self.send_error(404, "Not found")

    def log_message(self, format: str, *args: object) -> None:  # noqa: A003
        return


if __name__ == "__main__":
    refresh_workbook_data()
    server = ThreadingHTTPServer((HOST, PORT), RequestHandler)
    print(f"Servidor listo en http://{HOST}:{PORT}")
    print(f"Excel cargado desde: {EXCEL_PATH}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nServidor detenido.")
    finally:
        server.server_close()
